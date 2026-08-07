"""
Convierte las hojas 'BASE DATOS' y 'CLIENTES' del Excel original a un JSON
que el seed de Prisma puede consumir directamente.

Uso:
    python3 scripts/import_from_excel.py /ruta/al/Programa_Ventas.xlsm

Genera: prisma/seed-data.json

Nota: se descartan filas sin código (producto) o sin RIF/CI (cliente),
que en el archivo original son celdas de plantilla/fórmulas vacías.
"""
import json
import sys
from pathlib import Path

import openpyxl

ITEM_TYPE_MAP = {
    "MP": "MP",
    "PPRINC": "PPrinc",
    "PSECUN": "PSecun",
    "PFINAL": "PFinal",
}


def clean_str(value) -> str:
    return str(value).strip() if value is not None else ""


def main() -> None:
    if len(sys.argv) < 2:
        print("Uso: python3 import_from_excel.py <archivo.xlsm>")
        sys.exit(1)

    source_path = Path(sys.argv[1])
    wb = openpyxl.load_workbook(source_path, data_only=True, keep_vba=True)

    products = []
    ws = wb["BASE DATOS"]
    for r in range(2, ws.max_row + 1):
        code = ws.cell(r, 1).value
        if not code:
            continue
        cost = ws.cell(r, 5).value
        products.append(
            {
                "code": clean_str(code),
                "description": clean_str(ws.cell(r, 2).value),
                "presentation": clean_str(ws.cell(r, 3).value) or "UN",
                "itemType": ITEM_TYPE_MAP.get(clean_str(ws.cell(r, 4).value).upper(), "PFinal"),
                "cost": float(cost) if isinstance(cost, (int, float)) else 0,
                "profitMargin": float(ws.cell(r, 6).value or 0),
                "price": float(ws.cell(r, 7).value or 0),
                "priceListNumber": int(ws.cell(r, 8).value or 1),
                "stock": float(ws.cell(r, 13).value or 0),
                "reorderPoint": float(ws.cell(r, 14).value or 0),
            }
        )

    clients = []
    ws = wb["CLIENTES"]
    for r in range(2, ws.max_row + 1):
        tax_id = ws.cell(r, 1).value
        if not tax_id:
            continue
        clients.append(
            {
                "taxId": clean_str(tax_id),
                "name": clean_str(ws.cell(r, 2).value) or "SIN NOMBRE",
                "address": clean_str(ws.cell(r, 3).value) or None,
                "phone": clean_str(ws.cell(r, 4).value) or None,
                "instagram": clean_str(ws.cell(r, 5).value) or None,
                "email": clean_str(ws.cell(r, 6).value) or None,
            }
        )

    out_path = Path(__file__).resolve().parent.parent / "prisma" / "seed-data.json"
    out_path.write_text(json.dumps({"products": products, "clients": clients}, indent=2, ensure_ascii=False))

    print(f"✅ {len(products)} productos y {len(clients)} clientes exportados a {out_path}")


if __name__ == "__main__":
    main()
